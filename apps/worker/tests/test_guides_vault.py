from pathlib import Path
from typing import Any

import docx
import numpy as np
import pytest

from brandlens_worker.guides import (
    Chunk,
    RuleCardDraft,
    RuleCardDrafts,
    Section,
    assign_codes,
    chunk_sections,
    clamp_pages,
    dedupe,
    parse_docx,
    parse_markdown,
    parse_pdf,
)
from brandlens_worker.jobs.ingest_guide import extract_drafts
from brandlens_worker.llm.base import Meta
from brandlens_worker.tracing import Tracer
from brandlens_worker.vault import fact_embed_text, plan_embeddings, rule_embed_text, text_hash

GUIDE_PDF = Path(__file__).resolve().parents[3] / "fixtures" / "paywise" / "guide.pdf"


def draft(**kw: Any) -> RuleCardDraft:
    base: dict[str, Any] = {
        "dimension": "claims",
        "title": "No guaranteed returns",
        "summary": "s",
        "body": "b",
        "severity": "blocker",
    }
    return RuleCardDraft.model_validate({**base, **kw})


# ------------------------------------------------------------------ parsing


def test_parse_guide_pdf_sections_and_pages() -> None:
    sections = parse_pdf(GUIDE_PDF.read_bytes())
    by_title = {s.title.split(" ")[0]: s for s in sections}
    assert sections[0].title == "PayWise Brand Guide" and sections[0].level == 1
    expected_pages = {
        "TV-02": 1,
        "TV-04": 1,
        "G-07": 2,
        "CL-03": 2,
        "VS-01": 3,
        "VS-02": 3,
        "VS-03": 3,
        "A-01": 3,
        "A-02": 4,
    }
    for code, page in expected_pages.items():
        assert by_title[code].page == page, code
        assert by_title[code].level == 3
    assert by_title["1."].level == 2  # "1. Voice and tone"
    # running header/footer are not body text; ligatures are folded
    assert not any("Page 1 of" in s.text or "Brand Guide · v3.2" in s.text for s in sections)
    assert "confident" in by_title["TV-02"].title
    assert "assured" in by_title["CL-03"].text


def test_parse_markdown_and_docx(tmp_path: Path) -> None:
    md = parse_markdown("# Guide\nintro\n## Tone\nBe warm.\nNo hype.")
    assert [(s.title, s.level, s.text) for s in md] == [
        ("Guide", 1, "intro"),
        ("Tone", 2, "Be warm.\nNo hype."),
    ]

    d = docx.Document()
    d.add_heading("Guide", level=0)
    d.add_heading("Tone", level=1)
    d.add_paragraph("Be warm.")
    path = tmp_path / "g.docx"
    d.save(str(path))
    sections = parse_docx(path.read_bytes())
    assert [(s.title, s.level, s.text) for s in sections] == [
        ("Guide", 1, ""),
        ("Tone", 1, "Be warm."),
    ]


def test_chunking_marks_pages_and_respects_limit() -> None:
    sections = [
        Section(title=f"S{i}", level=2, page=1 + i // 3, text="word " * 200) for i in range(9)
    ]
    chunks = chunk_sections(sections, max_tokens=600)
    assert len(chunks) > 1
    assert all(len(c.text) <= 600 * 4 for c in chunks)
    assert chunks[0].text.startswith("[page 1]")
    assert all(f"[page {p}]" in "".join(c.text for c in chunks) for p in (1, 2, 3))
    huge = chunk_sections(
        [Section(title="Big", level=2, page=4, text="A sentence here. " * 2000)], max_tokens=500
    )
    assert len(huge) > 1 and all(c.pages == [4] for c in huge)


def test_guide_pdf_chunks_fit_the_token_budget() -> None:
    chunks = chunk_sections(parse_pdf(GUIDE_PDF.read_bytes()))
    assert 2 <= len(chunks) <= 6
    assert all(len(c.text) <= 700 * 4 for c in chunks)


# ------------------------------------------------------------------ extraction


class FakeRouter:
    def __init__(self, outputs: list[RuleCardDrafts]) -> None:
        self.outputs = list(outputs)
        self.users: list[str] = []

    async def generate(
        self, system: str, user: str, schema: Any, *, prompt_version: str, **_: Any
    ) -> Any:
        self.users.append(user)
        assert prompt_version == "rule_extract@1"
        return self.outputs.pop(0), Meta(provider="fake", model="fake-1", latency_ms=1)


async def test_extract_drafts_one_call_per_chunk_and_clamps_pages() -> None:
    chunks = [
        Chunk(text="[page 1]\n## Tone", pages=[1]),
        Chunk(text="[page 2]\n## Claims", pages=[2, 3]),
    ]
    router = FakeRouter(
        [
            RuleCardDrafts(
                cards=[
                    draft(
                        dimension="tone", code_suggestion="TV-02", source_page=1, severity="warning"
                    )
                ]
            ),
            RuleCardDrafts(
                cards=[
                    draft(code_suggestion="CL-03", source_page=99),
                    draft(title="Disclaimer", source_page=3),
                ]
            ),
        ]
    )
    drafts = await extract_drafts(chunks, router, Tracer(None), guide_id="g")
    assert len(router.users) == 2
    assert all(u.startswith("<guide>\n") and u.endswith("\n</guide>") for u in router.users)
    assert [d.source_page for d in drafts] == [1, 2, 3]  # 99 isn't in the chunk → first page


def test_clamp_pages_keeps_valid_pages() -> None:
    chunk = Chunk(text="", pages=[2, 3])
    assert [
        d.source_page for d in clamp_pages([draft(source_page=3), draft(source_page=None)], chunk)
    ] == [3, 2]


def test_clamp_pages_tidies_titles_and_hyphens() -> None:
    [d] = clamp_pages(
        [draft(title="CL-03 No market\u2011linked guarantees", body="plain\u00a0text")],
        Chunk("", [1]),
    )
    assert d.title == "No market-linked guarantees"
    assert d.code_suggestion == "CL-03"  # recovered from the title
    assert d.body == "plain text"
    [d2] = clamp_pages([draft(title="Image warmth (VS-02)", dimension="visual")], Chunk("", [1]))
    assert (d2.title, d2.code_suggestion) == ("Image warmth", "VS-02")


def test_dedupe_drops_near_duplicates() -> None:
    drafts = [draft(title="A"), draft(title="A again"), draft(title="B")]
    vecs = np.array([[1.0, 0.0], [0.99, 0.141], [0.0, 1.0]], dtype=np.float32)
    vecs /= np.linalg.norm(vecs, axis=1, keepdims=True)
    kept = dedupe(drafts, lambda texts: vecs)
    assert [d.title for d in kept] == ["A", "B"]


def test_assign_codes_versions_and_numbering() -> None:
    existing = {"CL-03": 1, "TV-02": 2, "TV-04": 1}
    coded = assign_codes(
        [
            draft(code_suggestion="CL-03"),  # exists → v2 proposal
            draft(code_suggestion="cl-9"),  # new code, normalised
            draft(code_suggestion="CL-03"),  # duplicate in batch → next free number
            draft(dimension="tone", code_suggestion="X-1", severity="warning"),  # wrong prefix
            draft(dimension="access", code_suggestion=None, severity="suggestion"),
        ],
        existing,
    )
    assert [(c.code, c.version) for c in coded] == [
        ("CL-03", 2),
        ("CL-09", 1),
        ("CL-10", 1),
        ("TV-05", 1),
        ("A-01", 1),
    ]


# ------------------------------------------------------------------ embeddings plan


def test_plan_embeddings_selects_only_missing_or_changed() -> None:
    rule = {
        "id": "r1",
        "title": "T",
        "summary": "S",
        "body": "B",
        "examples": [{"kind": "do", "text": "x"}],
    }
    fact = {"id": "f1", "statement": "1% exit load", "value": {"exit_load_pct": 1}}
    fresh_rule = {
        **rule,
        "id": "r2",
        "has_embedding": True,
        "embedding_source_hash": text_hash(rule_embed_text(rule)),
    }
    stale_rule = {**rule, "id": "r3", "has_embedding": True, "embedding_source_hash": "old"}
    fresh_fact = {
        **fact,
        "id": "f2",
        "has_embedding": True,
        "embedding_source_hash": text_hash(fact_embed_text(fact)),
    }
    tasks = plan_embeddings(
        [{**rule, "has_embedding": False}, fresh_rule, stale_rule],
        [{**fact, "has_embedding": False}, fresh_fact],
    )
    assert [(t.table, t.id) for t in tasks] == [
        ("rule_cards", "r1"),
        ("rule_cards", "r3"),
        ("facts", "f1"),
    ]
    assert "Do: x" in tasks[0].text
    assert '"exit_load_pct": 1' in tasks[2].text


@pytest.mark.parametrize("value", [None, {}])
def test_fact_embed_text_handles_empty_value(value: Any) -> None:
    assert fact_embed_text({"statement": "S", "value": value}) == "S\n{}"
